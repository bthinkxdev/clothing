# admin_dashboard/utils.py
from typing import Any, Dict, List
from datetime import datetime, timedelta
from django.utils import timezone
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from django.db.models import QuerySet, F
import json


class DateRangeHelper:
    """Helper for date range calculations"""
    
    @staticmethod
    def get_date_range(period: str) -> tuple:
        """Get start and end dates for common periods"""
        end_date = timezone.now()
        
        periods = {
            'today': timedelta(days=0),
            'yesterday': timedelta(days=1),
            'last_7_days': timedelta(days=7),
            'last_30_days': timedelta(days=30),
            'this_month': None,
            'last_month': None,
            'this_year': None,
        }
        
        if period == 'today':
            start_date = end_date.replace(hour=0, minute=0, second=0, microsecond=0)
        elif period == 'yesterday':
            start_date = (end_date - timedelta(days=1)).replace(hour=0, minute=0, second=0, microsecond=0)
            end_date = start_date.replace(hour=23, minute=59, second=59)
        elif period == 'this_month':
            start_date = end_date.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        elif period == 'last_month':
            first_day_this_month = end_date.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
            end_date = first_day_this_month - timedelta(seconds=1)
            start_date = end_date.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        elif period == 'this_year':
            start_date = end_date.replace(month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
        elif period in periods:
            start_date = end_date - periods[period]
        else:
            start_date = end_date - timedelta(days=30)
        
        return start_date, end_date
    
    @staticmethod
    def parse_custom_date_range(start_str: str, end_str: str) -> tuple:
        """Parse custom date range strings"""
        try:
            start_date = datetime.strptime(start_str, '%Y-%m-%d')
            end_date = datetime.strptime(end_str, '%Y-%m-%d')
            
            # Make timezone aware
            start_date = timezone.make_aware(start_date.replace(hour=0, minute=0, second=0))
            end_date = timezone.make_aware(end_date.replace(hour=23, minute=59, second=59))
            
            return start_date, end_date
        except ValueError:
            # Return default if parsing fails
            return DateRangeHelper.get_date_range('last_30_days')


class PaginationHelper:
    """Helper for pagination"""
    
    @staticmethod
    def paginate(queryset: QuerySet, page: int, per_page: int = 25) -> Dict:
        """Paginate queryset and return dict with page info"""
        paginator = Paginator(queryset, per_page)
        
        try:
            page_obj = paginator.page(page)
        except PageNotAnInteger:
            page_obj = paginator.page(1)
        except EmptyPage:
            page_obj = paginator.page(paginator.num_pages)
        
        return {
            'items': page_obj.object_list,
            'page': page_obj.number,
            'total_pages': paginator.num_pages,
            'total_items': paginator.count,
            'has_previous': page_obj.has_previous(),
            'has_next': page_obj.has_next(),
            'previous_page': page_obj.previous_page_number() if page_obj.has_previous() else None,
            'next_page': page_obj.next_page_number() if page_obj.has_next() else None,
        }


class ChartDataFormatter:
    """Format data for chart.js"""
    
    @staticmethod
    def format_line_chart(data: List[Dict], x_key: str, y_keys: List[str], labels: List[str] = None) -> Dict:
        """Format data for line chart"""
        if not labels:
            labels = y_keys
        
        datasets = []
        for i, y_key in enumerate(y_keys):
            datasets.append({
                'label': labels[i] if i < len(labels) else y_key,
                'data': [item[y_key] for item in data],
                'borderColor': ChartDataFormatter._get_color(i),
                'backgroundColor': ChartDataFormatter._get_color(i, alpha=0.1),
                'tension': 0.4
            })
        
        return {
            'labels': [item[x_key] for item in data],
            'datasets': datasets
        }
    
    @staticmethod
    def format_bar_chart(data: List[Dict], x_key: str, y_key: str, label: str = 'Data') -> Dict:
        """Format data for bar chart"""
        return {
            'labels': [item[x_key] for item in data],
            'datasets': [{
                'label': label,
                'data': [item[y_key] for item in data],
                'backgroundColor': [ChartDataFormatter._get_color(i, alpha=0.7) for i in range(len(data))],
                'borderColor': [ChartDataFormatter._get_color(i) for i in range(len(data))],
                'borderWidth': 1
            }]
        }
    
    @staticmethod
    def format_pie_chart(data: List[Dict], label_key: str, value_key: str) -> Dict:
        """Format data for pie/doughnut chart"""
        return {
            'labels': [item[label_key] for item in data],
            'datasets': [{
                'data': [item[value_key] for item in data],
                'backgroundColor': [ChartDataFormatter._get_color(i, alpha=0.8) for i in range(len(data))],
                'borderColor': ['#fff'] * len(data),
                'borderWidth': 2
            }]
        }
    
    @staticmethod
    def _get_color(index: int, alpha: float = 1.0) -> str:
        """Get color from palette"""
        colors = [
            (99, 102, 241),    # Indigo
            (139, 92, 246),    # Purple
            (236, 72, 153),    # Pink
            (239, 68, 68),     # Red
            (249, 115, 22),    # Orange
            (234, 179, 8),     # Yellow
            (34, 197, 94),     # Green
            (6, 182, 212),     # Cyan
            (59, 130, 246),    # Blue
            (168, 85, 247),    # Violet
        ]
        
        r, g, b = colors[index % len(colors)]
        return f'rgba({r}, {g}, {b}, {alpha})'


class FilterHelper:
    """Helper for building filters"""
    
    @staticmethod
    def build_order_filters(request) -> Dict:
        """Build order filters from request"""
        filters = {}
        
        if status := request.GET.get('status'):
            filters['status'] = status
        
        if payment_method := request.GET.get('payment_method'):
            filters['payments__method'] = payment_method
            # filters['payments__status'] = 'success'
        
        if customer_id := request.GET.get('customer'):
            filters['user_id'] = customer_id
        
        if search := request.GET.get('search'):
            from django.db.models import Q
            # Will be used with Q objects in view
            filters['_search'] = search
        
        # Date range
        date_range = request.GET.get('date_range', 'last_30_days')
        if date_range == 'custom':
            start = request.GET.get('start_date')
            end = request.GET.get('end_date')
            if start and end:
                start_date, end_date = DateRangeHelper.parse_custom_date_range(start, end)
            else:
                start_date, end_date = DateRangeHelper.get_date_range('last_30_days')
        else:
            start_date, end_date = DateRangeHelper.get_date_range(date_range)
        
        filters['placed_at__range'] = [start_date, end_date]
        
        return filters
    
    @staticmethod
    def build_product_filters(request) -> Dict:
        """Build product filters from request"""
        filters = {}
        
        if category_id := request.GET.get('category'):
            filters['category_id'] = category_id
        
        if brand := request.GET.get('brand'):
            filters['brand__iexact'] = brand
        
        if is_active := request.GET.get('is_active'):
            filters['is_active'] = is_active == 'true'
        
        if search := request.GET.get('search'):
            filters['_search'] = search
        
        return filters


class ExportHelper:
    """Helper for data export"""
    
    @staticmethod
    def export_to_csv(data: List[Dict], headers: List[str], filename: str) -> str:
        """Export data to CSV"""
        import csv
        from io import StringIO
        
        output = StringIO()
        writer = csv.writer(output)
        
        # Write headers
        writer.writerow(headers)
        
        # Write data
        for row in data:
            writer.writerow([row.get(h, '') for h in headers])
        
        return output.getvalue()
    
    @staticmethod
    def export_to_excel(data: List[Dict], headers: List[str], filename: str):
        """Export data to Excel (requires openpyxl)"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
            from io import BytesIO
            
            wb = Workbook()
            ws = wb.active
            ws.title = "Export"
            
            # Style headers
            header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            
            # Write headers
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center')
            
            # Write data
            for row_idx, row_data in enumerate(data, 2):
                for col_idx, header in enumerate(headers, 1):
                    ws.cell(row=row_idx, column=col_idx, value=row_data.get(header, ''))
            
            # Adjust column widths
            for column in ws.columns:
                max_length = 0
                column = [cell for cell in column]
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column[0].column_letter].width = adjusted_width
            
            # Save to BytesIO
            output = BytesIO()
            wb.save(output)
            output.seek(0)
            
            return output
            
        except ImportError:
            raise ImportError("openpyxl is required for Excel export. Install with: pip install openpyxl")


class PermissionHelper:
    """Helper for permission checks"""
    
    @staticmethod
    def check_admin_permission(user) -> bool:
        """Check if user has admin permissions"""
        if not user.is_authenticated:
            return False
        
        return user.is_staff or user.role in ['admin', 'staff']
    
    @staticmethod
    def check_staff_permission(user, permission: str) -> bool:
        """Check specific staff permission"""
        if not PermissionHelper.check_admin_permission(user):
            return False
        
        # You can implement granular permissions here
        # For now, admin has all permissions
        if user.role == 'admin':
            return True
        
        # Staff permissions can be role-based
        staff_permissions = {
            'view_orders': True,
            'edit_orders': True,
            'view_products': True,
            'edit_products': False,
            'view_customers': True,
            'edit_customers': False,
            'view_reports': True,
        }
        
        return staff_permissions.get(permission, False)


class NotificationHelper:
    """Helper for admin notifications"""
    
    @staticmethod
    def get_pending_actions() -> Dict:
        """Get counts of items needing attention"""
        from app.models import Order, Review, Inventory
        
        pending_orders = Order.objects.filter(status='pending').count()
        pending_reviews = Review.objects.filter(approved=False).count()
        low_stock = Inventory.objects.filter(quantity__lte=F('low_stock_threshold')).count()
        
        return {
            'pending_orders': pending_orders,
            'pending_reviews': pending_reviews,
            'low_stock': low_stock,
            'total': pending_orders + pending_reviews + low_stock
        }


class StatisticsHelper:
    """Helper for statistical calculations"""
    
    @staticmethod
    def calculate_growth_rate(current: float, previous: float) -> float:
        """Calculate percentage growth rate"""
        if previous == 0:
            return 100.0 if current > 0 else 0.0
        return ((current - previous) / previous) * 100
    
    @staticmethod
    def calculate_conversion_rate(converted: int, total: int) -> float:
        """Calculate conversion rate percentage"""
        if total == 0:
            return 0.0
        return (converted / total) * 100
    
    @staticmethod
    def calculate_average_with_fallback(values: List[float], fallback: float = 0.0) -> float:
        """Calculate average with fallback for empty lists"""
        if not values:
            return fallback
        return sum(values) / len(values)

